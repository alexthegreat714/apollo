"""
xlsx_ingest.py - Excel XLSX Ingestion for Apollo RAG

Converts Excel spreadsheet data to text documents for RAG ingestion.
"""

import os
import sys
from pathlib import Path
from datetime import datetime
from typing import Dict, Any, List

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from apollo_schemas.validators import normalize_document, validate_document
from apollo_ingest.chunker import TextChunker
from apollo_ingest.rag_writer import RAGWriter


def read_xlsx(file_path: str, sheet_name: str = None) -> Dict[str, List[Dict[str, Any]]]:
    """
    Read Excel file into dictionary of sheet data.

    Args:
        file_path: Path to XLSX file
        sheet_name: Specific sheet to read (all sheets if None)

    Returns:
        Dictionary mapping sheet names to list of row dictionaries
    """
    try:
        import openpyxl
    except ImportError:
        raise ImportError("Please install openpyxl: pip install openpyxl")

    wb = openpyxl.load_workbook(file_path, data_only=True)
    sheets_data = {}

    sheet_names = [sheet_name] if sheet_name else wb.sheetnames

    for name in sheet_names:
        if name not in wb.sheetnames:
            continue

        ws = wb[name]
        rows = list(ws.iter_rows(values_only=True))

        if not rows:
            continue

        # First row as headers
        headers = [str(h) if h else f"Column_{i}" for i, h in enumerate(rows[0])]

        # Convert remaining rows to dictionaries
        data = []
        for row in rows[1:]:
            if any(cell is not None for cell in row):
                row_dict = {}
                for i, value in enumerate(row):
                    if i < len(headers):
                        row_dict[headers[i]] = value
                data.append(row_dict)

        if data:
            sheets_data[name] = data

    wb.close()
    return sheets_data


def sheets_to_text(sheets_data: Dict[str, List[Dict[str, Any]]]) -> str:
    """
    Convert sheet data to text format.

    Args:
        sheets_data: Dictionary of sheet data

    Returns:
        Text representation
    """
    sections = []

    for sheet_name, rows in sheets_data.items():
        if not rows:
            continue

        lines = [f"=== Sheet: {sheet_name} ===", ""]

        for i, row in enumerate(rows):
            parts = []
            for key, value in row.items():
                if value is not None and str(value).strip():
                    parts.append(f"{key}: {value}")
            if parts:
                lines.append(f"Row {i+1}: " + ". ".join(parts) + ".")

        sections.append("\n".join(lines))

    return "\n\n".join(sections)


def infer_metadata(file_path: str, sheets_data: Dict[str, List[Dict[str, Any]]]) -> Dict[str, Any]:
    """
    Infer metadata from Excel file and content.
    """
    path = Path(file_path)
    stat = path.stat()
    modified = datetime.fromtimestamp(stat.st_mtime)

    filename_lower = path.stem.lower()

    # Count total rows
    total_rows = sum(len(rows) for rows in sheets_data.values())

    # Get all headers
    all_headers = set()
    for rows in sheets_data.values():
        if rows:
            all_headers.update(rows[0].keys())

    # Infer kind
    if any(word in filename_lower for word in ['statement', 'balance', 'income', 'cash']):
        kind = 'statement'
    elif any(word in filename_lower for word in ['forecast', 'projection', 'budget']):
        kind = 'projection'
    elif any(word in filename_lower for word in ['report', 'analysis']):
        kind = 'report'
    else:
        kind = 'report'

    return {
        'kind': kind,
        'source': path.stem,
        'timestamp': modified.isoformat(),
        'file_path': str(path.absolute()),
        'file_type': 'xlsx',
        'file_size': stat.st_size,
        'sheet_count': len(sheets_data),
        'total_rows': total_rows,
        'sheets': list(sheets_data.keys()),
        'columns': list(all_headers),
    }


def ingest_xlsx(
    file_path: str,
    sheet_name: str = None,
    kind: str = None,
    source: str = None,
    chunk_size: int = 512,
    chunk_overlap: int = 50,
    rag_dir: str = None,
    collection: str = None
) -> Dict[str, Any]:
    """
    Ingest an Excel XLSX file into Apollo RAG.

    Args:
        file_path: Path to XLSX file
        sheet_name: Specific sheet to ingest (all if None)
        kind: Document kind
        source: Source name
        chunk_size: Size of text chunks
        chunk_overlap: Overlap between chunks
        rag_dir: ChromaDB persist directory
        collection: Collection name

    Returns:
        Ingestion result dictionary
    """
    path = Path(file_path)

    if not path.exists():
        return {"error": f"File not found: {file_path}"}

    if path.suffix.lower() not in ['.xlsx', '.xls']:
        return {"error": f"Not an Excel file: {file_path}"}

    print(f"[XLSX] Reading: {path.name}")

    # Read Excel
    try:
        sheets_data = read_xlsx(file_path, sheet_name)
    except Exception as e:
        return {"error": f"Failed to read Excel: {e}"}

    if not sheets_data:
        return {"error": "Excel file contains no data"}

    # Convert to text
    text = sheets_to_text(sheets_data)

    if not text or len(text.strip()) < 50:
        return {"error": "Excel contains insufficient content"}

    # Infer metadata
    meta = infer_metadata(file_path, sheets_data)
    if kind:
        meta['kind'] = kind
    if source:
        meta['source'] = source

    # Create summary
    summary = f"Excel workbook with {meta['sheet_count']} sheets and {meta['total_rows']} total rows. Sheets: {', '.join(meta['sheets'][:3])}"
    if len(meta['sheets']) > 3:
        summary += f" and {len(meta['sheets']) - 3} more"

    # Build document
    doc = {
        'kind': meta['kind'],
        'timestamp': meta['timestamp'],
        'source': meta['source'],
        'summary': summary,
        'full_text': text,
        'metadata': {
            'file_path': meta['file_path'],
            'file_type': meta['file_type'],
            'sheet_count': meta['sheet_count'],
            'total_rows': meta['total_rows'],
            'sheets': ','.join(meta['sheets']),
        }
    }

    # Normalize and validate
    try:
        doc = normalize_document(doc, meta['kind'])
        is_valid, errors = validate_document(doc)
        if not is_valid:
            print(f"[XLSX] Validation warnings: {errors}")
    except Exception as e:
        print(f"[XLSX] Validation error: {e}")

    # Chunk the text
    chunker = TextChunker(chunk_size=chunk_size, chunk_overlap=chunk_overlap)
    chunks = chunker.chunk(text, {
        'kind': meta['kind'],
        'source': meta['source'],
        'file_path': meta['file_path'],
    })

    print(f"[XLSX] Created {len(chunks)} chunks from {meta['total_rows']} rows")

    # Write to RAG
    writer = RAGWriter(persist_dir=rag_dir, collection_name=collection)
    result = writer.write_chunks(chunks, base_id=doc.get('id', path.stem))

    result['file'] = str(path.name)
    result['kind'] = meta['kind']
    result['text_length'] = len(text)
    result['chunks'] = len(chunks)
    result['sheets'] = meta['sheet_count']
    result['rows'] = meta['total_rows']

    print(f"[XLSX] Ingestion complete: {result['count']} documents written")

    return result


if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser(description="Ingest Excel XLSX into Apollo RAG")
    parser.add_argument("file", help="XLSX file to ingest")
    parser.add_argument("--sheet", help="Specific sheet name")
    parser.add_argument("--kind", help="Document kind")
    parser.add_argument("--source", help="Source name")
    parser.add_argument("--chunk-size", type=int, default=512)
    parser.add_argument("--chunk-overlap", type=int, default=50)
    parser.add_argument("--rag-dir", help="ChromaDB directory")
    parser.add_argument("--collection", help="Collection name")

    args = parser.parse_args()

    result = ingest_xlsx(
        args.file,
        sheet_name=args.sheet,
        kind=args.kind,
        source=args.source,
        chunk_size=args.chunk_size,
        chunk_overlap=args.chunk_overlap,
        rag_dir=args.rag_dir,
        collection=args.collection
    )

    print("\nResult:")
    for key, value in result.items():
        print(f"  {key}: {value}")
